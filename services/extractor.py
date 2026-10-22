"""Text extraction for uploaded files, reporting progress through a callback.

progress(percent_within_stage: int, message: str)
"""

import base64
import csv
import io
import logging
from typing import Callable, Optional

from config import Config
from services import ollama_client

log = logging.getLogger(__name__)

Progress = Callable[[int, str], None]
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}

try:
    import fitz  # PyMuPDF
except Exception:  # pragma: no cover
    fitz = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None

_paddle = None


class ExtractionError(RuntimeError):
    pass


class ExtractionResult:
    def __init__(self, text: str, pages: Optional[int] = None, ocr_used: bool = False, ocr_engine: str = ""):
        self.text = text
        self.pages = pages
        self.ocr_used = ocr_used
        self.ocr_engine = ocr_engine


# ---------------------------------------------------------------- plain text
def decode_text(data: bytes) -> str:
    if data.startswith((b"\x00\x00\xfe\xff", b"\xff\xfe\x00\x00")):
        return data.decode("utf-32")
    if data.startswith((b"\xfe\xff", b"\xff\xfe")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _csv_to_text(data: bytes) -> str:
    decoded = decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(decoded, newline=""), dialect)
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows if any(c.strip() for c in row))


# ---------------------------------------------------------------- OCR engines
def _ocr_paddle(image_bytes: bytes) -> str:
    global _paddle
    import numpy as np
    import cv2
    from paddleocr import PaddleOCR

    if _paddle is None:
        _paddle = PaddleOCR(lang="en", use_doc_orientation_classify=True, use_doc_unwarping=False,
                            use_textline_orientation=False, device="cpu", enable_mkldnn=False)
    image = cv2.imdecode(np.frombuffer(image_bytes, np.uint8), cv2.IMREAD_COLOR)
    if image is None:
        raise ExtractionError("Image could not be decoded")
    result = _paddle.predict(image)
    lines = []
    for item in result or []:
        data = item if isinstance(item, dict) else getattr(item, "json", {}) or {}
        data = data.get("res", data)
        lines.extend(str(t) for t in data.get("rec_texts", []) if t)
    return "\n".join(lines)


def _ocr_tesseract(image_bytes: bytes) -> str:
    import pytesseract
    from PIL import Image

    if Config.TESSERACT_CMD:
        pytesseract.pytesseract.tesseract_cmd = Config.TESSERACT_CMD
    image = Image.open(io.BytesIO(image_bytes))
    return pytesseract.image_to_string(image)


def _ocr_ollama(image_bytes: bytes) -> str:
    if not Config.OLLAMA_VISION_MODEL:
        raise ExtractionError("OLLAMA_VISION_MODEL is not set")
    return ollama_client.vision_ocr(base64.b64encode(image_bytes).decode(), Config.OLLAMA_VISION_MODEL)


_ENGINES = {"paddle": _ocr_paddle, "tesseract": _ocr_tesseract, "ollama": _ocr_ollama}


def ocr_image(image_bytes: bytes) -> tuple[str, str]:
    """Run OCR with the configured engine (or first working one in auto mode)."""
    choice = Config.OCR_ENGINE
    if choice == "off":
        raise ExtractionError("OCR is disabled (OCR_ENGINE=off), so images/scanned pages cannot be read.")
    order = [choice] if choice in _ENGINES else ["paddle", "tesseract", "ollama"]
    errors = []
    for name in order:
        try:
            text = _ENGINES[name](image_bytes)
            return text.strip(), name
        except Exception as exc:  # engine missing or failed; try the next one
            log.info("OCR engine %s unavailable: %s", name, exc)
            errors.append(f"{name}: {exc}")
    raise ExtractionError(
        "No OCR engine is available. Install Tesseract (set TESSERACT_CMD), PaddleOCR, "
        "or set OLLAMA_VISION_MODEL. Details: " + " | ".join(errors)
    )


# ---------------------------------------------------------------- formats
def _pdf(data: bytes, progress: Progress) -> ExtractionResult:
    if fitz is None:
        raise ExtractionError("PyMuPDF is not installed")
    chunks, ocr_used, engine = [], False, ""
    with fitz.open(stream=data, filetype="pdf") as doc:
        total = len(doc)
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text.strip():
                progress(int(i / total * 100), f"Reading page {i} of {total}")
                chunks.append(text)
            else:
                progress(int(i / total * 100), f"Scanning page {i} of {total} with OCR")
                png = page.get_pixmap(dpi=200, alpha=False).tobytes("png")
                ocr_text, engine = ocr_image(png)
                ocr_used = True
                chunks.append(ocr_text)
    return ExtractionResult("\n\n".join(chunks), pages=total, ocr_used=ocr_used, ocr_engine=engine)


def _docx(data: bytes, progress: Progress) -> ExtractionResult:
    if DocxDocument is None:
        raise ExtractionError("python-docx is not installed")
    doc = DocxDocument(io.BytesIO(data))
    parts = []
    paragraphs = doc.paragraphs
    for i, para in enumerate(paragraphs, start=1):
        text = para.text.strip()
        if text:
            style = (para.style.name or "").lower() if para.style is not None else ""
            parts.append(f"## {text}" if style.startswith("heading") else text)
        if i % 50 == 0:
            progress(int(i / max(1, len(paragraphs)) * 80), f"Reading paragraph {i} of {len(paragraphs)}")
    for t_index, table in enumerate(doc.tables, start=1):
        progress(85, f"Reading table {t_index} of {len(doc.tables)}")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    progress(100, "Document read")
    return ExtractionResult("\n".join(parts))


def extract(data: bytes, ext: str, progress: Progress) -> ExtractionResult:
    ext = ext.lower()
    if ext in {".txt", ".md"}:
        progress(100, "Reading text file")
        return ExtractionResult(decode_text(data))
    if ext == ".csv":
        progress(100, "Parsing CSV rows")
        return ExtractionResult(_csv_to_text(data))
    if ext == ".pdf":
        return _pdf(data, progress)
    if ext == ".docx":
        return _docx(data, progress)
    if ext in IMAGE_EXTS:
        progress(30, "Running OCR on image")
        text, engine = ocr_image(data)
        progress(100, f"OCR complete ({engine})")
        return ExtractionResult(text, pages=1, ocr_used=True, ocr_engine=engine)
    raise ExtractionError(f"Unsupported file type: {ext}")
